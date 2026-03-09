# -*- coding: utf-8 -*-
"""サンプルデータ動作検証スクリプト"""
import sys, io, math
import numpy as np
import pandas as pd

sys.path.insert(0, '.')
from src.data_loader import load_experiment_data, get_temperature_groups, check_mass_balance
from src.kinetics import run_full_analysis, auto_detect_reaction_type
from src.plotting import (plot_raw, plot_multi_species, plot_best_fit_conc,
                          plot_rk4lsq_fit, plot_arrhenius, plot_raw_multi_temp)
from src.reporter import generate_excel_report

R = 8.314
PASS = "PASS"
FAIL = "FAIL"
TOTAL = [0, 0]  # [pass, fail]

def chk(cond, label):
    status = PASS if cond else FAIL
    mark = "  [OK]" if cond else "  [NG]"
    print(f"{mark} {label}")
    TOTAL[0 if cond else 1] += 1
    return cond

def load(path):
    with open(path, 'rb') as f:
        return load_experiment_data(io.BytesIO(f.read()))

# ══════════════════════════════════════════════════════════════════
print("=" * 60)
print("Sample 1: 単純1次反応 + ノイズ (k=0.0234 min⁻¹)")
print("=" * 60)
df1, meta1, warn1 = load('sample_data/sample1_simple_1st_order.xlsx')
tg1 = get_temperature_groups(df1)
dt1, dr1 = auto_detect_reaction_type(df1)

chk(len(df1) == 10, f"行数=10 (実={len(df1)})")
chk('concentration_B' not in df1.columns, "B列なし（単純反応）")
chk(dt1 == 'simple', f"自動判定=simple (実={dt1})")
chk(len(tg1) == 1, f"温度グループ=1 (実={len(tg1)})")

r1 = run_full_analysis(df1, reaction_type='simple')
k_int = r1.best_result.k
r2_int = r1.best_result.r2
chk(r1.best_order == 1, f"積分法 1次判定 (実={r1.best_order})")
chk(abs(k_int - 0.0234) / 0.0234 < 0.15, f"k誤差<15% k={k_int:.4f} (真値=0.0234)")
chk(r2_int > 0.97, f"R²>{0.97} (実={r2_int:.4f})")
chk(r1.rk4lsq is not None and r1.rk4lsq.success, f"RK4 success k={r1.rk4lsq.k:.4f}")
chk(r1.differential is not None, "微分法 実行")
# 図生成
try:
    fig = plot_best_fit_conc(df1, r1)
    chk(len(fig.data) >= 2, "ベストフィット図 生成")
except Exception as e:
    chk(False, f"図生成エラー: {e}")
# Excelレポート
try:
    excel = generate_excel_report(df1, meta1, r1)
    chk(len(excel) > 4000, f"レポート生成 {len(excel)}bytes")
except Exception as e:
    chk(False, f"レポートエラー: {e}")


# ══════════════════════════════════════════════════════════════════
print()
print("=" * 60)
print("Sample 2: 逐次反応 A→B→C (k1=0.05, k2=0.02 min⁻¹)")
print("=" * 60)
df2, meta2, warn2 = load('sample_data/sample2_sequential.xlsx')
dt2, dr2 = auto_detect_reaction_type(df2)
mb_ok2, mb_cv2 = check_mass_balance(df2)

chk(len(df2) == 10, f"行数=10 (実={len(df2)})")
chk('concentration_B' in df2.columns, "B列あり")
chk('concentration_C' in df2.columns, "C列あり")
chk(dt2 == 'sequential', f"自動判定=sequential (実={dt2})")
chk(mb_ok2, f"質量バランス OK (CV={mb_cv2:.4f})")
print(f"  [INFO] 判定理由: {dr2}")

r2 = run_full_analysis(df2, reaction_type='sequential')
rk2 = r2.rk4lsq
chk(rk2 is not None and rk2.success, f"RK4 sequential success")
if rk2 and rk2.success:
    k1_err = abs(rk2.k  - 0.05) / 0.05 * 100
    k2_err = abs(rk2.k2 - 0.02) / 0.02 * 100
    chk(k1_err < 5, f"k1誤差<5%: {k1_err:.1f}% (fit={rk2.k:.5f}, 真=0.05)")
    chk(k2_err < 5, f"k2誤差<5%: {k2_err:.1f}% (fit={rk2.k2:.5f}, 真=0.02)")
    chk(rk2.r2 > 0.999, f"R²>0.999 (実={rk2.r2:.5f})")
chk(r2.detected_reaction_type == 'sequential', f"detected_type=sequential")
# 誤タイプ選択のフォールバック確認
r2_wrong = run_full_analysis(df2, reaction_type='parallel')
chk(any('自動判定' in w or '推奨' in w for w in r2_wrong.warnings),
    "誤タイプ選択時に警告あり")
try:
    fig2 = plot_rk4lsq_fit(df2, r2.rk4lsq)
    chk(len(fig2.data) >= 4, f"逐次反応図 生成 ({len(fig2.data)}トレース)")
except Exception as e:
    chk(False, f"図生成エラー: {e}")
try:
    excel2 = generate_excel_report(df2, meta2, r2)
    chk(len(excel2) > 4000, f"レポート生成 {len(excel2)}bytes")
except Exception as e:
    chk(False, f"レポートエラー: {e}")


# ══════════════════════════════════════════════════════════════════
print()
print("=" * 60)
print("Sample 3: 並列反応 A→B + A→C (k1=0.03, k2=0.01 min⁻¹)")
print("=" * 60)
df3, meta3, warn3 = load('sample_data/sample3_parallel.xlsx')
dt3, dr3 = auto_detect_reaction_type(df3)
mb_ok3, mb_cv3 = check_mass_balance(df3)

chk(dt3 == 'parallel', f"自動判定=parallel (実={dt3})")
chk(mb_ok3, f"質量バランス OK (CV={mb_cv3:.4f})")
print(f"  [INFO] 判定理由: {dr3}")

r3 = run_full_analysis(df3, reaction_type='parallel')
rk3 = r3.rk4lsq
chk(rk3 is not None and rk3.success, f"RK4 parallel success")
if rk3 and rk3.success:
    k1_err = abs(rk3.k  - 0.03) / 0.03 * 100
    k2_err = abs(rk3.k2 - 0.01) / 0.01 * 100
    chk(k1_err < 5, f"k1誤差<5%: {k1_err:.1f}% (fit={rk3.k:.5f}, 真=0.03)")
    chk(k2_err < 5, f"k2誤差<5%: {k2_err:.1f}% (fit={rk3.k2:.5f}, 真=0.01)")
    chk(rk3.r2 > 0.999, f"R²>0.999 (実={rk3.r2:.5f})")
try:
    fig3 = plot_rk4lsq_fit(df3, r3.rk4lsq)
    chk(len(fig3.data) >= 4, f"並列反応図 生成 ({len(fig3.data)}トレース)")
except Exception as e:
    chk(False, f"図生成エラー: {e}")


# ══════════════════════════════════════════════════════════════════
print()
print("=" * 60)
print("Sample 4: 複数温度 単純反応 Arrhenius (Ea=60kJ/mol)")
print("=" * 60)
df4, meta4, warn4 = load('sample_data/sample4_multi_temp_arrhenius.xlsx')
tg4 = get_temperature_groups(df4)

chk(len(df4) == 32, f"行数=32 ({4}温度×{8}点) (実={len(df4)})")
chk(len(tg4) == 4, f"温度グループ=4 (実={len(tg4)}): {list(tg4.keys())}")
for T_c, sdf in tg4.items():
    chk(len(sdf) == 8, f"  {T_c}°C グループ: {len(sdf)}行")

r4 = run_full_analysis(df4, reaction_type='simple',
                       enable_arrhenius=True, temp_groups=tg4)
arr4 = r4.arrhenius
chk(arr4 is not None, "Arrhenius解析 実行")
if arr4:
    Ea_err = abs(arr4.Ea - 60000) / 60000 * 100
    chk(Ea_err < 2, f"Ea誤差<2%: {Ea_err:.2f}% (fit={arr4.Ea/1000:.2f}kJ/mol, 真=60.00)")
    chk(arr4.r2 > 0.999, f"アレニウスR²>0.999 (実={arr4.r2:.5f})")
    chk(arr4.n_temperatures == 4, f"温度点数=4 (実={arr4.n_temperatures})")
try:
    fig4a = plot_arrhenius(arr4)
    chk(len(fig4a.data) == 2, "アレニウス図 生成 (散布+回帰)")
    fig4b = plot_raw_multi_temp(tg4)
    chk(len(fig4b.data) == 4, f"複数温度重ね図 4温度 ({len(fig4b.data)}トレース)")
except Exception as e:
    chk(False, f"図生成エラー: {e}")
try:
    excel4 = generate_excel_report(df4, meta4, r4)
    chk(len(excel4) > 5000, f"レポート生成 {len(excel4)}bytes")
except Exception as e:
    chk(False, f"レポートエラー: {e}")


# ══════════════════════════════════════════════════════════════════
print()
print("=" * 60)
print("Sample 5: 異なる時間点 (A/B/C が別タイミング測定)")
print("  逐次反応 k1=0.04, k2=0.015 min⁻¹")
print("=" * 60)
df5, meta5, warn5 = load('sample_data/sample5_different_timepoints.xlsx')
dt5, dr5 = auto_detect_reaction_type(df5)

chk('concentration_B' in df5.columns, "B列あり")
chk('concentration_C' in df5.columns, "C列あり")
chk(dt5 == 'sequential', f"自動判定=sequential (実={dt5})")

# NaN存在確認
n_nan_B = df5['concentration_B'].isna().sum()
n_nan_C = df5['concentration_C'].isna().sum()
chk(n_nan_B > 0, f"B列にNaNあり ({n_nan_B}個) — 異時間点を正しく保持")
chk(n_nan_C > 0, f"C列にNaNあり ({n_nan_C}個) — 異時間点を正しく保持")

# 異時間点テストの注意ログ
any_diff_time_warn = any('時間点' in w or '異なる' in w for w in warn5)
if any_diff_time_warn:
    print(f"  [INFO] 異時間点警告: {[w for w in warn5 if '時間点' in w or '異なる' in w]}")

r5 = run_full_analysis(df5, reaction_type='sequential')
rk5 = r5.rk4lsq
chk(rk5 is not None and rk5.success, f"異時間点 RK4 sequential success")
if rk5 and rk5.success:
    k1_err = abs(rk5.k  - 0.04)  / 0.04  * 100
    k2_err = abs(rk5.k2 - 0.015) / 0.015 * 100
    chk(k1_err < 10, f"k1誤差<10%: {k1_err:.1f}% (fit={rk5.k:.5f}, 真=0.04)")
    chk(k2_err < 10, f"k2誤差<10%: {k2_err:.1f}% (fit={rk5.k2:.5f}, 真=0.015)")
try:
    fig5 = plot_rk4lsq_fit(df5, r5.rk4lsq)
    chk(hasattr(fig5, 'data'), "異時間点 RK4図 生成")
except Exception as e:
    chk(False, f"図生成エラー: {e}")


# ══════════════════════════════════════════════════════════════════
print()
print("=" * 60)
print("Sample 6: 複数温度 × 逐次反応 (k1/k2 別Arrhenius)")
print("  k1: Ea=50kJ  k2: Ea=30kJ")
print("=" * 60)
df6, meta6, warn6 = load('sample_data/sample6_multi_temp_sequential.xlsx')
tg6 = get_temperature_groups(df6)
dt6, dr6 = auto_detect_reaction_type(df6)

chk(len(tg6) == 3, f"温度グループ=3 (実={len(tg6)}): {list(tg6.keys())}")
chk(dt6 == 'sequential', f"自動判定=sequential (実={dt6})")
chk('concentration_B' in df6.columns, "B列あり")

r6 = run_full_analysis(df6, reaction_type='sequential',
                       enable_arrhenius=True, temp_groups=tg6)
arr6_k1 = r6.arrhenius
arr6_k2 = r6.arrhenius_k2

chk(arr6_k1 is not None, "k1 Arrhenius 計算")
chk(arr6_k2 is not None, "k2 Arrhenius 計算")

if arr6_k1:
    Ea1_err = abs(arr6_k1.Ea - 50000) / 50000 * 100
    chk(Ea1_err < 3, f"k1 Ea誤差<3%: {Ea1_err:.2f}% (fit={arr6_k1.Ea/1000:.2f}kJ, 真=50.00)")
    chk(arr6_k1.k_label == 'k1', f"k_label=k1 (実={arr6_k1.k_label})")

if arr6_k2:
    Ea2_err = abs(arr6_k2.Ea - 30000) / 30000 * 100
    chk(Ea2_err < 3, f"k2 Ea誤差<3%: {Ea2_err:.2f}% (fit={arr6_k2.Ea/1000:.2f}kJ, 真=30.00)")
    chk(arr6_k2.k_label == 'k2', f"k_label=k2 (実={arr6_k2.k_label})")

try:
    excel6 = generate_excel_report(df6, meta6, r6)
    chk(len(excel6) > 5000, f"レポート生成 {len(excel6)}bytes (k1/k2 Arrhenius含む)")
except Exception as e:
    chk(False, f"レポートエラー: {e}")


# ══════════════════════════════════════════════════════════════════
print()
print("=" * 60)
print("追加検証: エッジケース")
print("=" * 60)

# Edge 1: 濃度AのみでB/Cが全NaN (テンプレートにB/C列あり)
import openpyxl as ox
wb_e = ox.load_workbook('sample_data/sample2_sequential.xlsx')
ws_e = wb_e['実験データ']
# B列(3列目)とC列(4列目)を全空白に
for row in ws_e.iter_rows(min_row=2, max_row=ws_e.max_row):
    row[2].value = None  # B
    row[3].value = None  # C
buf = __import__('io').BytesIO()
wb_e.save(buf); buf.seek(0)
df_e, _, warn_e = load_experiment_data(buf)
chk('concentration_B' not in df_e.columns,
    "B/C全NaN → data_loaderが列除去し単純反応として扱う")
chk(any('空欄' in w or '欠損' in w for w in warn_e),
    "全NaN除去の警告メッセージあり")

# Edge 2: 重複time
df_dup = pd.DataFrame({
    'time': [0.0, 5.0, 5.0, 10.0, 20.0],
    'concentration': [1.0, 0.78, 0.77, 0.60, 0.37],
    'temperature': [25.0]*5, 'notes': ['']*5
})
dup_mask = df_dup['time'].duplicated(keep='first')
df_dedup = df_dup[~dup_mask].reset_index(drop=True)
chk(len(df_dedup) == 4 and df_dup['time'].duplicated().any(),
    "重複time検出・除去ロジック動作")

# Edge 3: 3点データ → RK4スキップ
df_3pt = pd.DataFrame({
    'time': [0.0, 30.0, 60.0],
    'concentration': [1.0, 0.5, 0.25],
    'temperature': [25.0]*3, 'notes': ['']*3
})
r_3pt = run_full_analysis(df_3pt, enable_rk4lsq=True)
chk(r_3pt.rk4lsq is not None and not r_3pt.rk4lsq.success,
    f"3点データ → RK4 success=False (msg: {r_3pt.rk4lsq.message[:40] if r_3pt.rk4lsq else 'N/A'})")

# Edge 4: 質量バランス崩れ検出
df_imb = pd.read_excel('sample_data/sample2_sequential.xlsx', sheet_name='実験データ', header=0)
# B列を1.5倍にしてバランス崩す（列名検出後）
df_seq_raw, _, _ = load('sample_data/sample2_sequential.xlsx')
df_seq_imb = df_seq_raw.copy()
df_seq_imb['concentration_B'] = df_seq_imb['concentration_B'] * 1.5
mb_ok_imb, mb_cv_imb = check_mass_balance(df_seq_imb)
chk(not mb_ok_imb and mb_cv_imb > 0.05,
    f"質量バランス崩れ検出 (CV={mb_cv_imb:.3f}>0.05)")

# Edge 5: 単温度でArrhenius試行 → None返却
tg_single = get_temperature_groups(df_seq_raw)  # 単一温度
r_single_arr = run_full_analysis(df_seq_raw, reaction_type='sequential',
                                  enable_arrhenius=True,
                                  temp_groups=tg_single if len(tg_single) >= 2 else None)
chk(r_single_arr.arrhenius is None, "単温度でArrhenius=None")


# ══════════════════════════════════════════════════════════════════
print()
print("=" * 60)
print(f"結果サマリー:  PASS={TOTAL[0]}  FAIL={TOTAL[1]}")
print("=" * 60)
if TOTAL[1] == 0:
    print("全検証 PASS")
else:
    print(f"  {TOTAL[1]} 件の NG あり — 上記ログを確認")

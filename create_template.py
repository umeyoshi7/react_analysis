"""Script to create the experiment Excel template."""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "template", "experiment_template.xlsx")

# Sample data: sequential reaction A→B→C with k1=0.05, k2=0.02 min⁻¹ (approx)
import math
SAMPLE_TIMES = [0, 5, 10, 15, 20, 30, 45, 60, 90, 120]
# A: C_A = exp(-k1*t), k1=0.05
SAMPLE_CONC_A = [round(math.exp(-0.05 * t), 4) for t in SAMPLE_TIMES]
# B: analytic for A→B→C
def _B(t, k1=0.05, k2=0.02):
    return k1 / (k2 - k1) * (math.exp(-k1 * t) - math.exp(-k2 * t))
SAMPLE_CONC_B = [round(_B(t), 4) for t in SAMPLE_TIMES]
# C: mass balance C = 1 - A - B
SAMPLE_CONC_C = [round(max(1.0 - SAMPLE_CONC_A[i] - SAMPLE_CONC_B[i], 0.0), 4) for i in range(len(SAMPLE_TIMES))]
SAMPLE_TEMP  = [25.0] * len(SAMPLE_TIMES)
SAMPLE_NOTES = ["開始"] + [""] * (len(SAMPLE_TIMES) - 2) + ["終了"]


def thin_border():
    side = Side(style="thin")
    return Border(left=side, right=side, top=side, bottom=side)


def header_fill():
    return PatternFill("solid", fgColor="4472C4")


def create_template():
    wb = openpyxl.Workbook()

    # ---- Sheet 1: 実験データ ----
    ws1 = wb.active
    ws1.title = "実験データ"

    headers = [
        "時間 (Time)",
        "濃度_A (Concentration_A)",
        "濃度_B (Concentration_B)",
        "濃度_C (Concentration_C)",
        "温度 (Temperature)",
        "備考 (Notes)",
    ]
    units = ["min", "mol/L", "mol/L", "mol/L", "°C", "-"]

    # Header row
    for col, (h, u) in enumerate(zip(headers, units), start=1):
        cell = ws1.cell(row=1, column=col, value=f"{h}\n({u})")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill()
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = thin_border()

    ws1.row_dimensions[1].height = 32

    # Data rows
    for i, (t, cA, cB, cC, temp, note) in enumerate(
        zip(SAMPLE_TIMES, SAMPLE_CONC_A, SAMPLE_CONC_B, SAMPLE_CONC_C, SAMPLE_TEMP, SAMPLE_NOTES),
        start=2
    ):
        ws1.cell(row=i, column=1, value=t).border = thin_border()
        ws1.cell(row=i, column=2, value=cA).border = thin_border()
        ws1.cell(row=i, column=3, value=cB).border = thin_border()
        ws1.cell(row=i, column=4, value=cC).border = thin_border()
        ws1.cell(row=i, column=5, value=temp).border = thin_border()
        ws1.cell(row=i, column=6, value=note).border = thin_border()

    # Column widths
    ws1.column_dimensions["A"].width = 18
    ws1.column_dimensions["B"].width = 24
    ws1.column_dimensions["C"].width = 24
    ws1.column_dimensions["D"].width = 24
    ws1.column_dimensions["E"].width = 20
    ws1.column_dimensions["F"].width = 16

    # ---- Sheet 2: 実験条件 ----
    ws2 = wb.create_sheet("実験条件")

    condition_data = [
        ("ラベル", "値", "単位"),
        ("実験名", "逐次反応実験-001", ""),
        ("反応物質", "物質A→B→C", ""),
        ("初期濃度", 1.000, "mol/L"),
        ("反応温度", 25.0, "°C"),
        ("実験日", "2026-03-08", "date"),
        ("担当者", "山田太郎", ""),
        ("備考", "窒素雰囲気下", ""),
    ]

    for row_idx, (label, value, unit) in enumerate(condition_data, start=1):
        for col_idx, val in enumerate([label, value, unit], start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border()
            if row_idx == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill()
                cell.alignment = Alignment(horizontal="center")

    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 22
    ws2.column_dimensions["C"].width = 12

    os.makedirs(os.path.dirname(TEMPLATE_PATH), exist_ok=True)
    wb.save(TEMPLATE_PATH)
    print(f"Template saved: {TEMPLATE_PATH}")


if __name__ == "__main__":
    create_template()
